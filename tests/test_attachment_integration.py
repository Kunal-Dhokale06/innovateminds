import pytest
import io
import base64
import httpx
from pypdf import PdfWriter
from docx import Document

@pytest.fixture
def sample_pdf_bytes():
    pdf_writer = PdfWriter()
    pdf_writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    pdf_writer.write(buf)
    return buf.getvalue()

@pytest.fixture
def sample_docx_bytes():
    doc = Document()
    doc.add_paragraph("ModelRouter Enterprise Specification: Automated LLM Decisioning")
    doc.add_paragraph("Security Level: Confidential Tier 1")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

@pytest.fixture
def sample_png_bytes():
    # Valid 10x10 red PNG with full header and IEND chunk
    return base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAoAAAAKCAIAAAACUFjqAAAAE0lEQVR4nGP8z4APMOGVZRip0gBBLAETee26JgAAAABJRU5ErkJggg==")

@pytest.mark.asyncio
async def test_extract_attachments_endpoint(sample_pdf_bytes, sample_docx_bytes, sample_png_bytes):
    async with httpx.AsyncClient(base_url="http://127.0.0.1:8000") as client:
        files = [
            ("files", ("invoice.pdf", sample_pdf_bytes, "application/pdf")),
            ("files", ("spec.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("notes.txt", b"Deployment checklist: 1. Deploy backend. 2. Verify tests.", "text/plain")),
            ("files", ("ui.png", sample_png_bytes, "image/png"))
        ]
        resp = await client.post("/v1/attachments/extract", files=files)
        assert resp.status_code == 200
        data = resp.json()
        assert len(data["attachments"]) == 4
        assert data["has_images"] is True
        assert data["has_documents"] is True
        assert data["total_tokens"] > 0

        # Verify docx extracted text
        docx_item = next(a for a in data["attachments"] if a["name"] == "spec.docx")
        assert "Automated LLM Decisioning" in docx_item["preview_text"]

@pytest.mark.asyncio
async def test_document_injected_in_completion(sample_docx_bytes):
    b64_doc = base64.b64encode(sample_docx_bytes).decode("utf-8")
    payload = {
        "model": "auto",
        "messages": [{"role": "user", "content": "Analyze the attached specification."}],
        "attachments": [
            {
                "name": "spec.docx",
                "type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "data": b64_doc
            }
        ]
    }
    async with httpx.AsyncClient(base_url="http://127.0.0.1:8000") as client:
        resp = await client.post("/v1/chat/completions", json=payload)
        assert resp.status_code == 200
        assert resp.headers.get("x-modelrouter-attachments-count") == "1"
        data = resp.json()
        assert "_modelrouter_attachments" in data
        assert data["_modelrouter_attachments"][0]["name"] == "spec.docx"

@pytest.mark.asyncio
async def test_image_requires_vision_and_disqualifies_text_only_models(sample_png_bytes):
    b64_png = f"data:image/png;base64,{base64.b64encode(sample_png_bytes).decode('utf-8')}"
    payload = {
        "model": "auto",
        "messages": [{"role": "user", "content": "Describe what is shown in this image."}],
        "attachments": [
            {
                "name": "diagram.png",
                "type": "image/png",
                "data": b64_png
            }
        ]
    }
    async with httpx.AsyncClient(base_url="http://127.0.0.1:8000") as client:
        resp = await client.post("/v1/chat/completions", headers={"x-modelrouter-explain": "true"}, json=payload)
        assert resp.status_code == 200
        assert resp.headers.get("x-modelrouter-requires-vision") == "true"
        
        # Verify selected model supports vision!
        selected_model = resp.headers.get("x-modelrouter-selected-model")
        # Must be one of the vision-enabled models
        assert selected_model in ("mock-flagship", "gemini-1.5-flash", "gemini-1.5-pro", "gpt-4o", "gpt-4o-mini")

        # Verify candidate evaluation matrix rejected text-only models with MISSING_FEATURE
        data = resp.json()
        evals = data["_modelrouter_decision"]["candidate_evaluations"]
        text_only_eval = next(e for e in evals if e["model_name"] == "mock-balanced")
        assert not text_only_eval["is_eligible"]
        assert any("Candidate does not support required multimodal vision / image inputs" in r for r in text_only_eval["rejection_reasons"])

@pytest.mark.asyncio
async def test_empty_attachment_error():
    async with httpx.AsyncClient(base_url="http://127.0.0.1:8000") as client:
        files = [("files", ("empty.txt", b"", "text/plain"))]
        resp = await client.post("/v1/attachments/extract", files=files)
        assert resp.status_code == 400
        assert "EMPTY_FILE" in resp.text

@pytest.mark.asyncio
async def test_unsupported_media_type_error():
    async with httpx.AsyncClient(base_url="http://127.0.0.1:8000") as client:
        files = [("files", ("binary.exe", b"\x4d\x5a\x90\x00", "application/octet-stream"))]
        resp = await client.post("/v1/attachments/extract", files=files)
        assert resp.status_code == 400
        assert "UNSUPPORTED_MEDIA_TYPE" in resp.text

@pytest.mark.asyncio
async def test_oversized_file_error():
    # 11 MB > 10 MB limit
    oversized = b"A" * (11 * 1024 * 1024)
    async with httpx.AsyncClient(base_url="http://127.0.0.1:8000") as client:
        files = [("files", ("huge.txt", oversized, "text/plain"))]
        resp = await client.post("/v1/attachments/extract", files=files)
        assert resp.status_code == 400
        assert "FILE_TOO_LARGE" in resp.text
