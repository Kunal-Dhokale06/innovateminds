import io
import os
import re
import base64
import math
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass, field
from fastapi import HTTPException
from PIL import Image

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024     # 10 MB per file
MAX_TOTAL_SIZE_BYTES = 25 * 1024 * 1024    # 25 MB total
MAX_ATTACHMENTS_COUNT = 5

SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}
SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
ALL_SUPPORTED_EXTENSIONS = SUPPORTED_DOCUMENT_EXTENSIONS | SUPPORTED_IMAGE_EXTENSIONS

MIME_TO_EXT = {
    "application/pdf": ".pdf",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/webp": ".webp",
}

@dataclass
class ProcessedAttachment:
    filename: str
    media_category: str  # "document" or "image"
    mime_type: str
    size_bytes: int
    extracted_text: Optional[str] = None
    base64_data_url: Optional[str] = None
    token_estimate: int = 0
    image_width: Optional[int] = None
    image_height: Optional[int] = None

@dataclass
class AttachmentProcessingResult:
    attachments: List[ProcessedAttachment] = field(default_factory=list)
    has_images: bool = False
    has_documents: bool = False
    total_token_estimate: int = 0
    document_context_block: str = ""
    error_message: Optional[str] = None

class AttachmentProcessor:
    @staticmethod
    def _sanitize_filename(filename: str) -> str:
        # Strip path traversal characters and control characters
        cleaned = os.path.basename(filename)
        cleaned = re.sub(r'[^\w\.\-\s_]', '', cleaned).strip()
        return cleaned or "attachment"

    @staticmethod
    def _extract_text_from_pdf(content_bytes: bytes, filename: str) -> str:
        if not pypdf:
            raise HTTPException(status_code=500, detail="pypdf library not available for PDF extraction.")
        try:
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            if reader.is_encrypted:
                try:
                    reader.decrypt('')
                except Exception:
                    raise HTTPException(status_code=400, detail=f"PDF_ENCRYPTED: File '{filename}' is password protected.")
            
            extracted = []
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted.append(f"[Page {page_idx + 1}]\n{page_text.strip()}")
            
            full_text = "\n\n".join(extracted)
            if not full_text.strip():
                return f"[Document '{filename}' contains {len(reader.pages)} page(s) with no extractable text or is image-only scan.]"
            return full_text
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"CORRUPTED_DOCUMENT: Could not parse PDF '{filename}': {str(e)}")

    @staticmethod
    def _extract_text_from_docx(content_bytes: bytes, filename: str) -> str:
        if not docx:
            raise HTTPException(status_code=500, detail="python-docx library not available for DOCX extraction.")
        try:
            doc = docx.Document(io.BytesIO(content_bytes))
            lines = []
            for p in doc.paragraphs:
                if p.text.strip():
                    lines.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        lines.append(row_text)
            full_text = "\n".join(lines)
            if not full_text.strip():
                return f"[Document '{filename}' is empty or contains no extractable text paragraphs.]"
            return full_text
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"CORRUPTED_DOCUMENT: Could not parse DOCX '{filename}': {str(e)}")

    @staticmethod
    def _extract_text_from_plain(content_bytes: bytes, filename: str) -> str:
        try:
            return content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return content_bytes.decode("latin-1")
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"CORRUPTED_DOCUMENT: Could not decode text file '{filename}': {str(e)}")

    @staticmethod
    def _process_image(content_bytes: bytes, filename: str, mime_type: str) -> Tuple[str, int, int, int]:
        try:
            img = Image.open(io.BytesIO(content_bytes))
            img.verify()  # Verify header and integrity
            # Reopen to read dimensions safely
            img = Image.open(io.BytesIO(content_bytes))
            width, height = img.size
            fmt = (img.format or "PNG").lower()
            if fmt == "jpeg":
                fmt = "jpg"
            
            # Construct standard base64 data URL
            b64_str = base64.b64encode(content_bytes).decode("utf-8")
            actual_mime = f"image/{fmt}" if fmt in ("png", "jpeg", "webp") else mime_type or "image/png"
            data_url = f"data:{actual_mime};base64,{b64_str}"
            
            # Conservative token estimate for high-resolution images
            # Tile estimation: base 85 tokens + up to 4 tiles * 170 tokens ~ 765 tokens
            tiles = max(1, math.ceil((width * height) / (512 * 512)))
            token_estimate = min(1200, 85 + (min(tiles, 4) * 170))

            return data_url, token_estimate, width, height
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"INVALID_IMAGE: File '{filename}' is not a valid or readable image: {str(e)}")

    @classmethod
    def process_raw_attachment(
        cls,
        filename: str,
        content_bytes: bytes,
        mime_type: Optional[str] = None
    ) -> ProcessedAttachment:
        clean_name = cls._sanitize_filename(filename)
        size_bytes = len(content_bytes)

        if size_bytes == 0:
            raise HTTPException(status_code=400, detail=f"EMPTY_FILE: Attachment '{clean_name}' contains zero bytes.")

        if size_bytes > MAX_FILE_SIZE_BYTES:
            max_mb = MAX_FILE_SIZE_BYTES / (1024 * 1024)
            actual_mb = size_bytes / (1024 * 1024)
            raise HTTPException(
                status_code=400,
                detail=f"FILE_TOO_LARGE: File '{clean_name}' ({actual_mb:.2f} MB) exceeds maximum allowed size ({max_mb:.0f} MB)."
            )

        # Detect extension
        ext = os.path.splitext(clean_name)[1].lower()
        if not ext and mime_type in MIME_TO_EXT:
            ext = MIME_TO_EXT[mime_type]

        if ext not in ALL_SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(ALL_SUPPORTED_EXTENSIONS))
            raise HTTPException(
                status_code=400,
                detail=f"UNSUPPORTED_MEDIA_TYPE: Extension '{ext}' is not supported. Supported: {supported}"
            )

        # Document processing
        if ext in SUPPORTED_DOCUMENT_EXTENSIONS:
            if ext == ".pdf":
                extracted = cls._extract_text_from_pdf(content_bytes, clean_name)
                mime = "application/pdf"
            elif ext == ".docx":
                extracted = cls._extract_text_from_docx(content_bytes, clean_name)
                mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif ext == ".md":
                extracted = cls._extract_text_from_plain(content_bytes, clean_name)
                mime = "text/markdown"
            else: # .txt
                extracted = cls._extract_text_from_plain(content_bytes, clean_name)
                mime = "text/plain"

            token_estimate = max(1, math.ceil((len(extracted) / 3.5) * 1.15))

            return ProcessedAttachment(
                filename=clean_name,
                media_category="document",
                mime_type=mime,
                size_bytes=size_bytes,
                extracted_text=extracted,
                token_estimate=token_estimate
            )

        # Image processing
        else:
            default_mime = f"image/{ext.replace('.', '')}"
            mime = mime_type or default_mime
            data_url, token_est, w, h = cls._process_image(content_bytes, clean_name, mime)

            return ProcessedAttachment(
                filename=clean_name,
                media_category="image",
                mime_type=mime,
                size_bytes=size_bytes,
                base64_data_url=data_url,
                token_estimate=token_est,
                image_width=w,
                image_height=h
            )

    @classmethod
    def process_attachment_payloads(
        cls,
        attachments_input: List[Dict[str, Any]]
    ) -> AttachmentProcessingResult:
        """Processes a list of attachment dictionaries (e.g. from JSON body with base64 data)."""
        if not attachments_input:
            return AttachmentProcessingResult()

        if len(attachments_input) > MAX_ATTACHMENTS_COUNT:
            raise HTTPException(
                status_code=400,
                detail=f"MAX_ATTACHMENTS_EXCEEDED: Maximum allowed attachments per request is {MAX_ATTACHMENTS_COUNT} (received {len(attachments_input)})."
            )

        processed: List[ProcessedAttachment] = []
        total_size = 0
        has_images = False
        has_documents = False
        total_tokens = 0
        doc_blocks: List[str] = []

        for item in attachments_input:
            filename = item.get("name") or item.get("filename") or "attachment"
            mime_type = item.get("type") or item.get("mime_type")
            data_field = item.get("data") or item.get("content") or item.get("base64")

            if not data_field:
                raise HTTPException(status_code=400, detail=f"MISSING_DATA: Attachment '{filename}' has no content data.")

            # Decode data
            if isinstance(data_field, str):
                if "," in data_field and data_field.startswith("data:"):
                    # Extract from data URI
                    header, b64_part = data_field.split(",", 1)
                    if not mime_type:
                        mime_match = re.search(r'data:([^;]+);', header)
                        if mime_match:
                            mime_type = mime_match.group(1)
                    raw_bytes = base64.b64decode(b64_part)
                else:
                    raw_bytes = base64.b64decode(data_field)
            elif isinstance(data_field, bytes):
                raw_bytes = data_field
            else:
                raise HTTPException(status_code=400, detail=f"INVALID_DATA: Attachment '{filename}' content format invalid.")

            total_size += len(raw_bytes)
            if total_size > MAX_TOTAL_SIZE_BYTES:
                max_mb = MAX_TOTAL_SIZE_BYTES / (1024 * 1024)
                raise HTTPException(
                    status_code=400,
                    detail=f"TOTAL_SIZE_EXCEEDED: Cumulative attachments exceed {max_mb:.0f} MB."
                )

            att = cls.process_raw_attachment(filename, raw_bytes, mime_type)
            processed.append(att)

            if att.media_category == "image":
                has_images = True
            elif att.media_category == "document":
                has_documents = True
                size_kb = att.size_bytes / 1024.0
                doc_blocks.append(
                    f"[ATTACHED DOCUMENT: {att.filename} ({size_kb:.1f} KB)]\n"
                    f"--- Extracted Content ---\n"
                    f"{att.extracted_text}\n"
                    f"--- End of Document ---"
                )

            total_tokens += att.token_estimate

        doc_context = "\n\n".join(doc_blocks) if doc_blocks else ""

        return AttachmentProcessingResult(
            attachments=processed,
            has_images=has_images,
            has_documents=has_documents,
            total_token_estimate=total_tokens,
            document_context_block=doc_context
        )

attachment_processor = AttachmentProcessor()
